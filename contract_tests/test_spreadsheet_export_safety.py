"""Focused regression tests for safe downloadable spreadsheet exports."""
from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from employee_movements.services.export_safety import append_safe_xlsx_row, safe_csv_row, safe_csv_value

ROOT = Path(__file__).resolve().parents[1]


def test_csv_formula_prefixes_are_neutralized_after_leading_whitespace():
    values = safe_csv_row(["=1+1", "  =HYPERLINK('https://example.invalid')", "+cmd", "-cmd", "@SUM(A1:A2)", "normal"])
    assert values == ["'=1+1", "'  =HYPERLINK('https://example.invalid')", "'+cmd", "'-cmd", "'@SUM(A1:A2)", "normal"]


def test_csv_leading_control_characters_are_neutralized_and_typed_values_preserved():
    assert safe_csv_value("\t=1+1") == "'\t=1+1"
    assert safe_csv_value("\r@SUM(A1:A2)") == "'\r@SUM(A1:A2)"
    assert safe_csv_value(123) == 123
    assert safe_csv_value(None) is None


def test_normal_xlsx_exports_store_formula_like_text_as_literal_strings():
    workbook = Workbook()
    worksheet = workbook.active
    append_safe_xlsx_row(worksheet, ["=1+1", "=HYPERLINK(\"https://example.invalid\")", 12])
    assert worksheet["A1"].value == "=1+1"  # An initially empty sheet starts at row 1.
    assert worksheet["A1"].data_type == "s"

    # Existing static headers must remain in row 1 and exported data starts at row 2.
    worksheet_with_header = workbook.create_sheet("with-header")
    worksheet_with_header.append(["Header"])
    append_safe_xlsx_row(worksheet_with_header, ["=2+2", "literal", 34])
    assert worksheet_with_header["A1"].value == "Header"
    assert worksheet_with_header["A2"].value == "=2+2"
    assert worksheet_with_header["A2"].data_type == "s"

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    loaded = load_workbook(buffer, data_only=False)
    assert loaded.active["A1"].value == "=1+1"
    assert loaded.active["A1"].data_type == "s"
    assert loaded.active["B1"].data_type == "s"
    assert loaded.active["C1"].value == 12
    assert loaded["with-header"]["A1"].value == "Header"
    assert loaded["with-header"]["A2"].data_type == "s"


def test_write_only_xlsx_exports_store_formula_like_text_as_literal_strings():
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("export")
    append_safe_xlsx_row(worksheet, ["=1+1", "user name", 12], write_only=True)
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    loaded = load_workbook(buffer, data_only=False, read_only=True)
    row = next(loaded.active.iter_rows(values_only=False))
    assert row[0].value == "=1+1"
    assert row[0].data_type == "s"
    assert row[1].data_type == "s"
    assert row[2].value == 12


def test_sensitive_exports_and_authenticated_pages_use_safe_output_helpers():
    report_source = (ROOT / "employee_movements" / "blueprints" / "reports.py").read_text(encoding="utf-8")
    mission_source = (ROOT / "employee_movements" / "services" / "mission_documents.py").read_text(encoding="utf-8")
    hooks_source = (ROOT / "employee_movements" / "hooks.py").read_text(encoding="utf-8")
    access_source = (ROOT / "employee_movements" / "access.py").read_text(encoding="utf-8")
    assert "append_safe_xlsx_row(" in report_source
    assert "safe_csv_row(" in report_source
    assert "append_safe_xlsx_row(" in mission_source
    assert "private_authenticated_response" in hooks_source
    assert "private, no-store, max-age=0" in hooks_source
    assert "return safe_local_return_to(target, fallback)" in access_source
