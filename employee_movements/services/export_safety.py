"""Spreadsheet-export hardening for untrusted personnel fields.

Excel cells whose value begins with ``=`` are interpreted as formulas by
openpyxl unless the cell is forced to a string. CSV has no cell-type metadata,
so formula-like text receives a leading apostrophe before it is serialized.
Numeric/date values are preserved as typed values in XLSX exports.
"""
from __future__ import annotations

from typing import Any, Iterable

# Besides ASCII whitespace, include common Unicode zero-width characters that
# can be used to disguise a formula prefix during review of exported data.
_LEADING_IGNORABLE = " \t\r\n\v\f\ufeff\u200b\u200c\u200d\u2060\u00a0"
_CSV_FORMULA_PREFIXES = ("=", "+", "-", "@")
_CSV_CONTROL_PREFIXES = ("\t", "\r", "\n")


def safe_csv_value(value: Any) -> Any:
    """Return a CSV-safe value without changing numeric/date objects.

    CSV has no mechanism to declare a field as text. An apostrophe prefix is
    therefore added when a text field begins with a formula trigger (including
    after whitespace/zero-width characters), or with a control character that
    spreadsheet software may trim before evaluating a formula.
    """
    if not isinstance(value, str) or not value:
        return value
    stripped = value.lstrip(_LEADING_IGNORABLE)
    if value[0] in _CSV_CONTROL_PREFIXES or (stripped and stripped[0] in _CSV_FORMULA_PREFIXES):
        return "'" + value
    return value


def safe_csv_row(values: Iterable[Any]) -> list[Any]:
    """Sanitize all fields in one CSV row while preserving non-text types."""
    return [safe_csv_value(value) for value in values]


def safe_xlsx_cell(worksheet, value: Any, *, write_only: bool = False, row: int | None = None, column: int | None = None):
    """Create a cell where every Python string is explicitly stored as text.

    Write-only worksheets require ``WriteOnlyCell`` objects. For ordinary
    worksheets, provide the destination row and column so the cell receives
    valid 1-based coordinates before it is serialized.
    """
    if write_only:
        from openpyxl.cell import WriteOnlyCell

        cell = WriteOnlyCell(worksheet, value=value)
    else:
        if row is None or column is None or row < 1 or column < 1:
            raise ValueError("row and column are required for ordinary worksheets")
        cell = worksheet.cell(row=row, column=column)
        cell.value = value
    if isinstance(value, str):
        # Prevent openpyxl's default formula inference for values beginning
        # with '='; this is especially important for user-controlled exports.
        cell.data_type = "s"
    return cell


def append_safe_xlsx_row(worksheet, values: Iterable[Any], *, write_only: bool = False) -> None:
    """Append a row while forcing user-controlled strings to literal text."""
    row_values = list(values)
    if write_only:
        worksheet.append([safe_xlsx_cell(worksheet, value, write_only=True) for value in row_values])
        return
    worksheet_is_empty = (
        worksheet.max_row == 1
        and worksheet.max_column == 1
        and worksheet.cell(row=1, column=1).value is None
    )
    target_row = 1 if worksheet_is_empty else worksheet.max_row + 1
    for column, value in enumerate(row_values, start=1):
        safe_xlsx_cell(worksheet, value, row=target_row, column=column)
