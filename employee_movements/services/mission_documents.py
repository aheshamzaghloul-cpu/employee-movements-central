"""Document generation services for assignment missions (مأموريات)."""

import io
import os
import tempfile
from datetime import datetime
from html import escape

from PIL import Image

from ..access import movement_is_closed
from .export_safety import append_safe_xlsx_row, safe_xlsx_cell

import pymupdf as fitz
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def mission_state_label(mission):
    # Keep display and lifecycle controls aligned with the shared legacy/current
    # closure rule; older records may have assignment_state='مغلق' only.
    if movement_is_closed(mission):
        return 'مغلقة'
    return getattr(mission, 'mission_state', None) or 'تحت التحرير'


def mission_template_pdf(
    mission,
    employee,
    branch,
    destination,
    root_path,
    creator=None,
    print_to_date=None,
    closer=None,
    ticket_image_data=None,
):
    """Render the supplied mission PDF template without changing its geometry."""
    template = os.path.join(root_path, 'static', 'mission', 'mission_template.pdf')
    font = os.path.join(root_path, 'static', 'mission', 'mission-original.ttf')
    doc = fitz.open(template)
    page = doc[0]

    variable_regions = [
        fitz.Rect(497.5, 74.0, 553.8, 88.8),
        fitz.Rect(458.0, 89.5, 553.8, 106.2),
        fitz.Rect(31.0, 80.0, 115.0, 94.0),
        fitz.Rect(62.0, 94.0, 84.5, 109.5),
        fitz.Rect(346.6, 158.7, 482.9, 177.2),
        fitz.Rect(156.8, 158.7, 274.4, 177.2),
        fitz.Rect(57.1, 158.7, 88.5, 177.2),
        fitz.Rect(346.5, 188.8, 482.9, 207.2),
        fitz.Rect(299.9, 253.9, 360.2, 272.6),
        fitz.Rect(428.8, 253.9, 483.1, 272.6),
        fitz.Rect(100.0, 281.5, 180.0, 298.5),
        fitz.Rect(100.0, 300.0, 180.0, 317.0),
    ]
    for rect in variable_regions:
        page.add_redact_annot(rect, fill=(1, 1, 1))
    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE)

    if os.path.exists(font):
        page.insert_font(fontfile=font, fontname='missionorig')
    else:
        fallback = os.path.join(root_path, 'static', 'mission', 'NotoNaskhArabic-Regular.ttf')
        if os.path.exists(fallback):
            page.insert_font(fontfile=fallback, fontname='missionorig')

    now = datetime.now()

    def put(rect, text, size=9.9603748, align='right', direction='rtl'):
        text = '' if text is None else str(text)
        safe = text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
        html = (
            f'<div style="font-family:missionorig;font-size:{size:.6f}pt;'
            f'line-height:1;white-space:nowrap;text-align:{align};direction:{direction};">{safe}</div>'
        )
        page.insert_htmlbox(rect, html)

    status = 'مغلقة' if mission_state_label(mission) == 'مغلقة' else 'تحت التحرير'
    put(fitz.Rect(498.0, 74.2, 552.8, 87.9), f'{mission.id} {status}', 8.0403004, 'right', 'ltr')

    # The existing top-right name slot shows exactly one person: the creator
    # while the mission is open, or the closer after it is closed. Keep the
    # template geometry and all other printed content unchanged.
    is_closed = mission_state_label(mission) == 'مغلقة'
    display_user = closer if is_closed else creator
    display_name = display_user.full_name if display_user else ''
    creator_job = creator.job_title if creator and creator.job_title else ''
    if display_name:
        put(fitz.Rect(462.5, 90.2, 552.8, 98.2), f'- {display_name}', 6.0002327, 'right', 'rtl')
    if creator_job:
        put(fitz.Rect(462.5, 97.3, 552.8, 106.0), creator_job, 6.0002327, 'right', 'rtl')

    put(fitz.Rect(32.0, 82.0, 77.8, 93.0), now.strftime('%Y/%m/%d'), 8.0403004, 'left', 'ltr')
    put(fitz.Rect(80.5, 82.0, 115.0, 93.0), now.strftime('%H:%M:%S'), 8.0403004, 'left', 'ltr')
    put(fitz.Rect(62.0, 96.5, 84.5, 107.5), '1 \\ 1', 8.0403004, 'center', 'ltr')

    employee_name = employee.full_name if employee else ''
    gov_name = branch.governorate.name if branch and branch.governorate else ''
    branch_name = branch.name if branch else ''
    destination_gov = destination.governorate.name if destination and destination.governorate else ''
    destination_name = destination.name if destination else ''

    put(fitz.Rect(347.0, 159.1, 482.7, 176.9), employee_name, 9.9603748, 'right', 'rtl')
    basic_branch = f'{gov_name} - {branch_name}' if gov_name and branch_name else gov_name or branch_name
    put(fitz.Rect(157.0, 159.1, 274.2, 176.9), basic_branch, 9.9603748, 'right', 'rtl')
    put(
        fitz.Rect(57.2, 159.1, 88.4, 176.9),
        employee.job_code if employee and employee.job_code else '',
        9.9603748,
        'center',
        'ltr',
    )

    mission_dest = (
        f'{destination_gov} - {destination_name}'
        if destination_gov and destination_name
        else destination_gov or destination_name
    )
    put(fitz.Rect(346.8, 189.2, 482.7, 206.9), mission_dest, 9.9603748, 'right', 'rtl')

    effective_to_date = print_to_date or mission.to_date
    put(
        fitz.Rect(300.1, 254.6, 359.9, 271.9),
        effective_to_date.strftime('%Y/%m/%d') if effective_to_date else '',
        9.9603748,
        'center',
        'ltr',
    )
    put(
        fitz.Rect(429.1, 254.6, 482.8, 271.9),
        mission.from_date.strftime('%Y/%m/%d') if mission.from_date else '',
        9.9603748,
        'center',
        'ltr',
    )

    put(fitz.Rect(116.8, 284.6, 172.2, 296.0), 'اعتماد مدير فرع', 9.9603748, 'center', 'rtl')
    put(fitz.Rect(107.9, 303.0, 170.1, 314.5), mission_dest, 9.9603748, 'center', 'rtl')

    # The original mission form remains page one with unchanged geometry. If a
    # ServiceDesk screenshot is linked to this mission, append it as a separate
    # evidence page so it is included in every official print/PDF path.
    if ticket_image_data:
        _append_ticket_image_page(doc, ticket_image_data, mission, employee, root_path)

    out = io.BytesIO()
    doc.save(out, garbage=4, deflate=True)
    doc.close()
    out.seek(0)
    return out.getvalue()


def _append_ticket_image_page(doc, ticket_image_data, mission, employee, root_path):
    """Append one A4 evidence page containing the saved ticket screenshot."""
    payload = bytes(ticket_image_data)
    with Image.open(io.BytesIO(payload)) as source:
        source.verify()
    with Image.open(io.BytesIO(payload)) as source:
        image_width, image_height = source.size

    page = doc.new_page(-1, width=595.32, height=841.92)
    font = os.path.join(root_path, 'static', 'mission', 'mission-original.ttf')
    if not os.path.exists(font):
        font = os.path.join(root_path, 'static', 'mission', 'NotoNaskhArabic-Regular.ttf')
    if os.path.exists(font):
        page.insert_font(fontfile=font, fontname='missionorig')

    employee_name = escape(employee.full_name if employee else '')
    title = escape(f'صورة تيكيت ServiceDesk — المأمورية رقم {mission.id}')
    details = f'<div style="font-family:missionorig;font-size:13pt;font-weight:bold;text-align:right;direction:rtl">{title}</div>'
    if employee_name:
        details += f'<div style="font-family:missionorig;font-size:9pt;text-align:right;direction:rtl">الموظف: {employee_name}</div>'
    page.insert_htmlbox(fitz.Rect(28, 15, 567, 58), details)

    box = fitz.Rect(28, 66, 567, 809)
    scale = min(box.width / max(1, image_width), box.height / max(1, image_height))
    width = image_width * scale
    height = image_height * scale
    image_rect = fitz.Rect(
        box.x0 + (box.width - width) / 2,
        box.y0 + (box.height - height) / 2,
        box.x0 + (box.width + width) / 2,
        box.y0 + (box.height + height) / 2,
    )
    page.insert_image(image_rect, stream=payload, keep_proportion=True)
    footer = '<div style="font-family:missionorig;font-size:8pt;text-align:center;direction:rtl;color:#526579">مرفق أصلي محفوظ مع سجل المأمورية</div>'
    page.insert_htmlbox(fitz.Rect(28, 814, 567, 837), footer)


def mission_export_xlsx(rows):
    """Build an XLSX export with a bounded sample and a streamed write-only workbook."""
    from itertools import islice

    wb = Workbook(write_only=True)
    ws = wb.create_sheet('المأموريات')
    # Keep one column for each distinct fact and align source/destination labels
    # with the all-results print report. The old workbook duplicated the source
    # governorate and branch under both “الأصلية” and “من” headings.
    headers = [
        'رقم المأمورية', 'كود شئون العاملين', 'اسم الموظف', 'محافظة الموظف',
        'فرع الموظف', 'محافظة جهة المأمورية', 'فرع جهة المأمورية',
        'من تاريخ', 'إلى تاريخ', 'عدد الأيام', 'حالة المأمورية', 'تاريخ الإنشاء',
        'تاريخ الإغلاق', 'منشئ المأمورية', 'سبب الإغلاق', 'ملاحظات',
    ]

    def _values(mission):
        employee = mission.employee
        origin = employee.branch if employee else None
        destination = mission.destination
        start = mission.from_date
        end = mission.to_date
        days = (end - start).days + 1 if start and end else None
        return [
            mission.id,
            employee.job_code if employee else '',
            employee.full_name if employee else '',
            origin.governorate.name if origin and origin.governorate else '',
            origin.name if origin else '',
            destination.governorate.name if destination and destination.governorate else '',
            destination.name if destination else '',
            start,
            end,
            days,
            mission_state_label(mission),
            mission.created_at,
            mission.closed_at,
            mission.creator_user.full_name if getattr(mission, 'creator_user', None) else '',
            mission.closure_reason or '',
            mission.notes or '',
        ]

    # Sample no more than 200 missions to keep practical column widths. This is
    # bounded memory and happens before the write-only worksheet starts writing.
    iterator = iter(rows)
    sample_missions = list(islice(iterator, 200))
    sample_values = [_values(mission) for mission in sample_missions]
    widths = [min(max(len(header) + 2, 12), 32) for header in headers]
    for values in sample_values:
        for index, value in enumerate(values):
            widths[index] = max(widths[index], min(max(len(str(value or '')) + 2, 12), 40))

    ws.freeze_panes = 'A2'
    for index, width in enumerate(widths, start=1):
        # In write-only mode, configure column dimensions before appending rows.
        ws.column_dimensions[get_column_letter(index)].width = width

    header_cells = []
    for header in headers:
        cell = safe_xlsx_cell(ws, header, write_only=True)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        header_cells.append(cell)
    ws.append(header_cells)

    exported_rows = 0
    for values in sample_values:
        append_safe_xlsx_row(ws, values, write_only=True)
        exported_rows += 1
    for mission in iterator:
        append_safe_xlsx_row(ws, _values(mission), write_only=True)
        exported_rows += 1

    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{exported_rows + 1}'
    # Keep small exports in memory, but let large workbooks spill to a temporary
    # file rather than duplicating the entire generated XLSX in process RAM.
    out = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode='w+b')
    wb.save(out)
    out.seek(0)
    return out

