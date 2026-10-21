"""Excel import wizard for governorates, branches, and employees."""

import io
import json
import logging
import os

from flask import Blueprint, abort, flash, redirect, render_template, request, session, url_for
from openpyxl import Workbook, load_workbook

from ..access import gids, has_role, log, req, roles
from ..extensions import db
from ..models import (
    Branch,
    BranchNameAlias,
    Employee,
    EmployeeNameAlias,
    Governorate,
    GovernorateNameAlias,
    TicketEmployeeCodeAlias,
)
from ..services.excel_import_matching import resolve_employee_import_branch
from ..services.entity_integrity import normalize_branch_code_key
from ..services.excel_import_apply import (
    apply_selected_excel_import_changes,
    validate_selected_excel_import_changes,
)
from ..services.excel_import_workflow import (
    build_excel_import_selection_plan,
    parse_excel_column_mapping,
    validate_excel_column_mapping,
)
from ..services.excel_import_schema import (
    _auto_map_headers,
    _cell_by_index,
    _excel_date,
    _excel_display_text,
    _excel_employee_code_key,
    _excel_employee_code_variants,
    _excel_mapped_cell_text,
    _excel_headers,
    _excel_issue_category,
    _excel_key,
    _excel_text,
    _field_labels,
    _import_spec,
    _validate_excel_dimensions,
)
from ..services.excel_import_storage import (
    _EXCEL_PENDING_TTL_SECONDS,
    _pending_excel_path,
    _pending_excel_review_path,
    _claim_pending_excel_apply,
    _release_pending_excel_apply,
    _remove_pending_excel,
    _save_pending_excel,
)
from ..services.private_work_files import (
    is_private_work_token,
    pending_files_expired,
    touch_private_files,
    write_private_bytes,
)
from ..validation import valid_email

bp = Blueprint('excel_import', __name__)
logger = logging.getLogger(__name__)


def _excel_import_permissions(kind):
    if kind not in ('governorates', 'branches', 'employees'):
        abort(404)
    if kind == 'governorates' and not has_role('مسؤول التطبيق'):
        abort(403)
    if kind == 'branches' and 'مسؤول التطبيق' not in roles():
        abort(403)
    if kind == 'employees' and not has_role('مسؤول التطبيق'):
        abort(403)


def _excel_scope_governorates():
    return (
        {g.id for g in Governorate.query.filter_by(is_active=True).all()}
        if 'مسؤول التطبيق' in roles()
        else set(gids())
    )


def _excel_import_page(kind, **ctx):
    _excel_import_permissions(kind)
    return render_template('excel_import.html', kind=kind, **ctx)


@bp.get('/excel-import/<kind>')
@req
def excel_import_page(kind):
    return _excel_import_page(kind)


@bp.get('/excel-import/<kind>/template')
@req
def excel_import_template(kind):
    _excel_import_permissions(kind)
    wb = Workbook()
    ws = wb.active
    ws.title = 'بيانات'
    if kind == 'governorates':
        headers = ['اسم المحافظة']
        ws.append(headers)
    elif kind == 'branches':
        headers = ['المحافظة', 'اسم الفرع', 'كود الفرع']
        ws.append(headers)
    else:
        headers = [
            'المحافظة',
            'الفرع',
            'كود الفرع',
            'اسم الموظف',
            'البريد الإلكتروني',
            'الوظيفة',
            'الكود الوظيفي',
            'تاريخ التعيين',
            'هاتف الشركة',
            'الهاتف الشخصي',
        ]
        ws.append(headers)
    for c in ws[1]:
        c.font = c.font.copy(bold=True)
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = max(
            16,
            min(32, max((len(_excel_text(c.value)) for c in col)) + 2),
        )
    data = io.BytesIO()
    wb.save(data)
    data.seek(0)
    from flask import send_file
    return send_file(
        data,
        as_attachment=True,
        download_name=f'{kind}-template.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


_EXCEL_IMPORT_KINDS = frozenset({'governorates', 'branches', 'employees'})


def _clear_excel_import_session():
    """Clear every session marker associated with the single active import wizard."""
    for key in ('excel_import_token', 'excel_review_token', 'excel_import_kind'):
        session.pop(key, None)


def _get_pending_excel(expected_kind=None):
    """Return the staged upload only for the import kind that created it.

    Import tokens are session-scoped, but the URL's ``kind`` is supplied by the
    client. Binding the token to the kind selected during preview prevents an
    old tab or edited request URL from reinterpreting one workbook as another
    entity type.
    """
    token = session.get('excel_import_token')
    stored_kind = session.get('excel_import_kind')
    if (
        not is_private_work_token(token)
        or stored_kind not in _EXCEL_IMPORT_KINDS
        or (expected_kind is not None and expected_kind != stored_kind)
    ):
        _remove_pending_excel(token)
        _remove_pending_excel(session.get('excel_review_token'))
        _clear_excel_import_session()
        return None, None
    path = _pending_excel_path(token)
    review_path = _pending_excel_review_path(token)
    if (
        not path
        or not os.path.isfile(path)
        or pending_files_expired((path, review_path), _EXCEL_PENDING_TTL_SECONDS)
    ):
        _remove_pending_excel(token)
        _remove_pending_excel(session.get('excel_review_token'))
        _clear_excel_import_session()
        return None, None
    touch_private_files((path, review_path))
    return token, path


def _render_mapping(kind, token, path, warning=None):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    headers = _excel_headers(ws)
    preview = []
    for row in ws.iter_rows(min_row=2, max_row=6, values_only=False):
        vals = [_excel_display_text(cell) for cell in row]
        if any(vals):
            preview.append(vals[:len(headers)])
    wb.close()
    spec = _import_spec(kind)
    mapping = _auto_map_headers(headers, spec)
    labels = _field_labels(kind)
    return _excel_import_page(
        kind,
        mapping_step=True,
        token=token,
        headers=headers,
        preview=preview,
        fields=labels,
        suggested=mapping,
        warning=warning,
    )


@bp.post('/excel-import/<kind>/preview')
@req
def excel_import_preview(kind):
    _excel_import_permissions(kind)
    f = request.files.get('file')
    raw_filename = getattr(f, 'filename', '') if f else ''
    filename = raw_filename.strip() if isinstance(raw_filename, str) else ''
    if not filename or not filename.lower().endswith(('.xlsx', '.xlsm', '.xltx')):
        flash('اختر ملف Excel بصيغة .xlsx أو .xlsm أو .xltx.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))
    token = None
    try:
        old_tokens = (session.get('excel_import_token'), session.get('excel_review_token'))
        for old_token in old_tokens:
            _remove_pending_excel(old_token)
        _clear_excel_import_session()
        (token, path) = _save_pending_excel(f)
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        _validate_excel_dimensions(ws)
        headers = _excel_headers(ws)
        wb.close()
        if not headers or not any(headers):
            _remove_pending_excel(token)
            flash('ملف Excel لا يحتوي على صف عناوين صالح.')
            return redirect(url_for('excel_import.excel_import_page', kind=kind))
        session['excel_import_token'] = token
        session['excel_import_kind'] = kind
        return _render_mapping(kind, token, path)
    except Exception as exc:
        logger.error('Excel import preview failed (kind=%s, error=%s)', kind, type(exc).__name__)
        # If the upload was persisted but parsing/mapping failed, remove both
        # temporary artifacts so repeated invalid uploads cannot accumulate.
        _remove_pending_excel(token)
        _clear_excel_import_session()
        flash('تعذر قراءة ملف Excel. تحقق من سلامة الملف وحجمه ثم حاول مرة أخرى.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))


def _mapped_value(row, mapping, field):
    return _excel_mapped_cell_text(row, mapping.get(field), field)


def _append_unique_catalog_candidate(catalog, key, record):
    """Keep every distinct candidate for a normalized catalogue label."""
    bucket = catalog.setdefault(key, [])
    record_id = getattr(record, 'id', id(record))
    if all(getattr(existing, 'id', id(existing)) != record_id for existing in bucket):
        bucket.append(record)


@bp.post('/excel-import/<kind>/confirm')
@req
def excel_import_confirm(kind):
    _excel_import_permissions(kind)
    (token, path) = _get_pending_excel(expected_kind=kind)
    if not token or not path:
        flash('انتهت جلسة مطابقة ملف Excel. ارفع الملف مرة أخرى.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))
    spec = _import_spec(kind)
    fields = _field_labels(kind)
    mapping = parse_excel_column_mapping(request.form, fields)
    mapping_warning = validate_excel_column_mapping(kind, spec, fields, mapping)
    if mapping_warning:
        return _render_mapping(kind, token, path, warning=mapping_warning)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        _validate_excel_dimensions(ws)
        headers = _excel_headers(ws)
        mapping_warning = validate_excel_column_mapping(
            kind, spec, fields, mapping, header_count=len(headers)
        )
        if mapping_warning:
            wb.close()
            return _render_mapping(kind, token, path, warning=mapping_warning)
        allowed_gids = _excel_scope_governorates()
        errors = []
        scoped_governorates = Governorate.query.filter(Governorate.is_active == True, Governorate.id.in_(allowed_gids)).all()
        # Keep every normalized canonical/alias candidate. A single-value dict
        # can silently bind an Excel row to whichever same-named governorate was
        # inserted last, which is unsafe for employee and branch imports.
        govs = {}
        for governorate in scoped_governorates:
            _append_unique_catalog_candidate(govs, _excel_key(governorate.name), governorate)
        scoped_governorates_by_id = {governorate.id: governorate for governorate in scoped_governorates}
        governorate_alias_keys_by_id = {governorate.id: set() for governorate in scoped_governorates}
        for alias in GovernorateNameAlias.query.filter(GovernorateNameAlias.governorate_id.in_(allowed_gids)).all():
            owner = scoped_governorates_by_id.get(alias.governorate_id)
            if owner:
                alias_key = _excel_key(alias.alias_text)
                _append_unique_catalog_candidate(govs, alias_key, owner)
                governorate_alias_keys_by_id.setdefault(owner.id, set()).add(alias_key)
        branches_by_gov = {}
        branch_codes_by_gov = {}
        branch_alias_keys_by_id = {}
        # Load the authorized branch catalog once. Name matching uses active
        # branches only, while code collision detection also sees inactive
        # records so an archived code cannot be silently reused.
        scoped_import_branches = (
            Branch.query.filter(Branch.governorate_id.in_(allowed_gids)).all()
            if allowed_gids else []
        )
        active_import_branches = [
            branch for branch in scoped_import_branches
            if branch.is_active and branch.governorate and branch.governorate.is_active
        ]
        for branch in scoped_import_branches:
            code_key = normalize_branch_code_key(branch.code)
            if code_key:
                branch_codes_by_gov.setdefault(branch.governorate_id, {}).setdefault(code_key, []).append(branch)
        for branch in active_import_branches:
            _append_unique_catalog_candidate(
                branches_by_gov.setdefault(branch.governorate_id, {}),
                _excel_key(branch.name), branch,
            )
            branch_alias_keys_by_id.setdefault(branch.id, set())
        active_branch_ids_for_import = [branch.id for branch in active_import_branches]
        if active_branch_ids_for_import:
            branch_by_id = {branch.id: branch for branch in active_import_branches}
            for alias in BranchNameAlias.query.filter(BranchNameAlias.branch_id.in_(active_branch_ids_for_import)).all():
                branch = branch_by_id.get(alias.branch_id)
                if branch:
                    _append_unique_catalog_candidate(
                        branches_by_gov.setdefault(branch.governorate_id, {}),
                        _excel_key(alias.alias_text), branch,
                    )
                    branch_alias_keys_by_id.setdefault(branch.id, set()).add(_excel_key(alias.alias_text))
        actions = []
        if kind == 'governorates':
            active_existing_governorates = Governorate.query.filter_by(is_active=True).all()
            existing = {_excel_key(g.name): g for g in active_existing_governorates}
            active_governorates_by_id = {governorate.id: governorate for governorate in active_existing_governorates}
            for alias in GovernorateNameAlias.query.all():
                owner = active_governorates_by_id.get(alias.governorate_id)
                if owner:
                    existing.setdefault(_excel_key(alias.alias_text), owner)
            pending = set(existing)
            for (rno, row) in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
                vals = {k: _mapped_value(row, mapping, k) for k in spec}
                name = vals.get('name', '')
                if not name:
                    continue
                key = _excel_key(name)
                if key in pending:
                    errors.append(f'صف {rno}: المحافظة موجودة بالفعل أو مكررة داخل الملف: {name}')
                    continue
                actions.append({
                    'type': 'new_governorate',
                    'row': rno,
                    'values': {'name': name},
                })
                pending.add(key)
        elif kind == 'branches':
            # Track targets/codes/names as we scan rows to prevent duplicate
            # actions within one workbook without an O(rows × actions) search.
            seen_branch_targets = set()
            pending_branch_codes = set()
            pending_branch_names = set()
            for (rno, row) in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
                vals = {k: _mapped_value(row, mapping, k) for k in spec}
                if not any(vals.values()):
                    continue
                governorate_candidates = govs.get(_excel_key(vals['governorate']), [])
                if len(governorate_candidates) > 1:
                    errors.append(
                        f"صف {rno}: اسم المحافظة أو اسمها البديل يطابق أكثر من محافظة؛ لا يمكن تحديد الهدف بالاسم وحده. صحح دليل المحافظات أو استخدم اسمًا فريدًا.",
                    )
                    continue
                governorate = governorate_candidates[0] if governorate_candidates else None
                if not governorate:
                    errors.append(
                        f"صف {rno}: المحافظة غير موجودة أو خارج نطاقك: {vals['governorate']}",
                    )
                    continue
                if not vals['name'] or not vals['code']:
                    errors.append(f'صف {rno}: اسم الفرع وكوده مطلوبان.')
                    continue

                code_key = normalize_branch_code_key(vals['code'])
                name_key = _excel_key(vals['name'])
                existing_codes = branch_codes_by_gov.get(governorate.id, {}).get(code_key, [])
                if len(existing_codes) > 1:
                    errors.append(
                        f"صف {rno}: كود الفرع {vals['code']} غير فريد داخل المحافظة؛ صحح الكود قبل الاستيراد.",
                    )
                    continue

                # The normalized branch code is authoritative. A changed or
                # misspelled branch name cannot cause the same code to create a
                # second branch or redirect the update to another record.
                branch = existing_codes[0] if existing_codes else None
                name_candidates = branches_by_gov.get(governorate.id, {}).get(name_key, [])
                if not branch and len(name_candidates) > 1:
                    errors.append(
                        f"صف {rno}: اسم الفرع «{vals['name']}» يطابق أكثر من سجل أو اسم بديل داخل المحافظة؛ لا يمكن تحديد الفرع بالاسم وحده. صحح الكود أو عالج التعارض في دليل الفروع."
                    )
                    continue
                if branch and (not branch.is_active or not branch.governorate or not branch.governorate.is_active):
                    errors.append(
                        f"صف {rno}: كود الفرع يخص سجلًا غير نشط؛ راجع حالة الفرع بدل إنشاء سجل مكرر.",
                    )
                    continue
                if not branch and len(name_candidates) == 1:
                    branch = name_candidates[0]

                if branch:
                    if branch.id in seen_branch_targets:
                        errors.append(f"صف {rno}: الفرع نفسه مذكور أكثر من مرة داخل الملف؛ احتفظ بصف واحد فقط.")
                        continue
                    # If a branch was found by name instead of code, only allow
                    # the proposed code update when that normalized code does
                    # not belong to any other active or archived branch.
                    conflicts = branch_codes_by_gov.get(governorate.id, {}).get(code_key, [])
                    if any(candidate.id != branch.id for candidate in conflicts):
                        errors.append(f"صف {rno}: كود الفرع {vals['code']} مستخدم بالفعل في فرع آخر.")
                        continue
                    changes = {}
                    is_known_alias = name_key in branch_alias_keys_by_id.get(branch.id, set())
                    name_owners = branches_by_gov.get(governorate.id, {}).get(name_key, [])
                    # Even if the label is an alias for this branch, it is not
                    # safe when the same normalized text identifies another
                    # branch's canonical name or alias in this governorate.
                    if any(owner.id != branch.id for owner in name_owners):
                        errors.append(f"صف {rno}: اسم الفرع أو اسمه البديل يطابق سجلًا آخر داخل المحافظة، بينما الكود يحدد فرعًا مختلفًا؛ عالج التعارض يدويًا.")
                        continue
                    pending_code_key = (governorate.id, code_key)
                    pending_name_key = (governorate.id, name_key)
                    if pending_code_key in pending_branch_codes and normalize_branch_code_key(branch.code) != code_key:
                        errors.append(f"صف {rno}: كود الفرع المقترح مستخدم في صف آخر من ملف Excel؛ راجع التعارض.")
                        continue
                    if vals['name'] and vals['name'] != branch.name and not is_known_alias:
                        if pending_name_key in pending_branch_names:
                            errors.append(f"صف {rno}: اسم الفرع المقترح مكرر في صف آخر من ملف Excel؛ راجع التعارض.")
                            continue
                        changes['name'] = {'old': branch.name, 'new': vals['name']}
                    if normalize_branch_code_key(vals['code']) != normalize_branch_code_key(branch.code):
                        changes['code'] = {'old': branch.code or '—', 'new': vals['code']}
                        pending_branch_codes.add(pending_code_key)
                    if 'name' in changes:
                        pending_branch_names.add(pending_name_key)
                    seen_branch_targets.add(branch.id)
                    if changes:
                        actions.append({
                            'type': 'branch',
                            'id': branch.id,
                            'row': rno,
                            'name': branch.name,
                            'governorate': governorate.name,
                            'changes': changes,
                        })
                    else:
                        actions.append({
                            'type': 'branch_match', 'id': branch.id, 'row': rno,
                            'name': branch.name, 'governorate': governorate.name,
                        })
                    continue

                pending_code_key = (governorate.id, code_key)
                pending_name_key = (governorate.id, name_key)
                if pending_code_key in pending_branch_codes:
                    errors.append(f"صف {rno}: كود الفرع مكرر داخل ملف Excel في المحافظة نفسها.")
                    continue
                if pending_name_key in pending_branch_names:
                    errors.append(f"صف {rno}: اسم الفرع مكرر داخل ملف Excel في المحافظة نفسها.")
                    continue
                # A code already present in any branch record is never reused,
                # including a record currently inactive or with a different name.
                if existing_codes:
                    errors.append(f"صف {rno}: كود الفرع {vals['code']} مستخدم بالفعل.")
                    continue
                actions.append({
                    'type': 'new_branch',
                    'row': rno,
                    'values': vals,
                    'governorate_id': governorate.id,
                    'governorate': governorate.name,
                })
                pending_branch_codes.add(pending_code_key)
                pending_branch_names.add(pending_name_key)
        else:
            # This importer is restricted to the application administrator.
            # Do not use bids() here: that helper represents the active
            # operational/work scope for ordinary users and can be narrowed
            # to the currently selected governorate. Employee administration
            # imports must resolve every active branch in every active
            # governorate the administrator manages, independent of the
            # current workspace selector.
            branch_cache = {}
            branches_by_code = {}
            branches_by_name = {}
            # Build code candidates from every branch record in the authorized
            # active-governorate catalogue, including archived branches. An
            # archived or duplicate code must not disappear from identity checks.
            for b in scoped_import_branches:
                if b.code not in (None, ''):
                    code_key = normalize_branch_code_key(b.code)
                    branches_by_code.setdefault(code_key, []).append(b)
            for b in active_import_branches:
                governorate_key = _excel_key(b.governorate.name)
                # Name-only matching accepts official names and confirmed aliases,
                # never a branch code silently masquerading as a branch name.
                canonical_key = _excel_key(b.name)
                _append_unique_catalog_candidate(branch_cache, (governorate_key, canonical_key), b)
                _append_unique_catalog_candidate(branches_by_name, canonical_key, b)
                for alias_key in branch_alias_keys_by_id.get(b.id, set()):
                    _append_unique_catalog_candidate(branch_cache, (governorate_key, alias_key), b)
                    _append_unique_catalog_candidate(branches_by_name, alias_key, b)

            # Build a lightweight identity index in one streaming database pass.
            # Retain only scalar identity fields, not every Employee ORM instance;
            # this keeps large Excel previews from filling memory with model graphs.
            code_buckets = {}
            email_buckets = {}
            employee_identity_by_id = {}
            employee_names_by_branch = {}
            identity_query = Employee.query.with_entities(
                Employee.id, Employee.branch_id, Employee.full_name,
                Employee.job_code, Employee.email, Employee.deleted_at,
            ).yield_per(2000)
            for employee_id, branch_id, full_name, job_code, email, deleted_at in identity_query:
                employee_identity_by_id[employee_id] = {
                    'branch_id': branch_id,
                    'full_name': full_name or '',
                    # Soft-deleted employees remain in the identity indexes so
                    # their codes/emails cannot be recycled as new people.
                    'is_deleted': deleted_at is not None,
                }
                code_key = _excel_employee_code_key(job_code)
                for variant in _excel_employee_code_variants(code_key):
                    code_buckets.setdefault(variant, set()).add(employee_id)
                email_key = _excel_key(email) if email else ''
                if email_key:
                    email_buckets.setdefault(email_key, set()).add(employee_id)
                name_key = _excel_key(full_name)
                if branch_id and name_key:
                    employee_names_by_branch.setdefault((branch_id, name_key), set()).add(employee_id)
            # Compare the confirmed alias text using the strict import normalizer.
            # The stored normalized_alias is optimized for OCR and may collapse O/0
            # or I/1; it must not become an unreviewed canonical-code match here.
            for alias_employee_id, alias_text in TicketEmployeeCodeAlias.query.with_entities(
                TicketEmployeeCodeAlias.employee_id, TicketEmployeeCodeAlias.alias_text
            ).yield_per(2000):
                for variant in _excel_employee_code_variants(alias_text):
                    code_buckets.setdefault(variant, set()).add(alias_employee_id)
            for alias_employee_id, alias_text in EmployeeNameAlias.query.with_entities(
                EmployeeNameAlias.employee_id, EmployeeNameAlias.alias_text
            ).yield_per(2000):
                owner = employee_identity_by_id.get(alias_employee_id)
                alias_key = _excel_key(alias_text)
                if owner and owner['branch_id'] and alias_key:
                    employee_names_by_branch.setdefault((owner['branch_id'], alias_key), set()).add(alias_employee_id)
            job_codes = {
                key: next(iter(ids)) for key, ids in code_buckets.items() if len(ids) == 1
            }
            ambiguous_job_codes = {key for key, ids in code_buckets.items() if len(ids) > 1}
            emails = {key: next(iter(ids)) for key, ids in email_buckets.items() if len(ids) == 1}
            ambiguous_emails = {key for key, ids in email_buckets.items() if len(ids) > 1}
            seen_job_codes = set()
            seen_emails = set()
            matched_employee_rows = {}
            pending_employee_names = set()
            for (rno, row) in enumerate(ws.iter_rows(min_row=2, values_only=False), 2):
                vals = {k: _mapped_value(row, mapping, k) for k in spec}
                if not any(vals.values()):
                    continue
                b, branch_match_method = resolve_employee_import_branch(
                    vals.get('branch_code', ''),
                    vals.get('governorate', ''),
                    vals.get('branch', ''),
                    branches_by_code=branches_by_code,
                    branches_by_name=branches_by_name,
                    govs=govs,
                    branch_cache=branch_cache,
                    governorate_alias_keys_by_id=governorate_alias_keys_by_id,
                    branch_alias_keys_by_id=branch_alias_keys_by_id,
                    normalize_key=_excel_key,
                    normalize_code_key=normalize_branch_code_key,
                )
                if b and branch_match_method == 'inactive_code':
                    errors.append(
                        f"صف {rno}: كود الفرع {vals.get('branch_code', '')} يعود إلى فرع غير نشط أو محافظة غير نشطة؛ لا يمكن ربط الموظف به. راجع سجل الفرع بدل استخدام الكود تلقائيًا."
                    )
                    continue
                if b and branch_match_method == 'code_label_mismatch':
                    errors.append(
                        f"صف {rno}: تعارض يحتاج مراجعة: كود الفرع {vals.get('branch_code', '')} يحدد «{b.name}» بمحافظة «{b.governorate.name if b.governorate else '—'}»، بينما الملف يذكر «{vals.get('branch', '')}» بمحافظة «{vals.get('governorate', '')}». صحح الكود أو الأسماء قبل الاستيراد."
                    )
                    continue
                if not b:
                    if branch_match_method == 'unknown_code':
                        errors.append(
                            f"صف {rno}: تعارض يحتاج مراجعة: كود الفرع {vals.get('branch_code', '')} غير موجود ضمن الفروع النشطة؛ صحح الكود. لم تتم المطابقة بالاسم لتجنب ربط الموظف بفرع غير مقصود."
                        )
                    elif branch_match_method == 'ambiguous_code':
                        errors.append(
                            f"صف {rno}: كود الفرع {vals.get('branch_code', '')} غير فريد؛ أضف اسم المحافظة الصحيح أو صحح كود الفرع."
                        )
                    elif branch_match_method == 'ambiguous_name':
                        errors.append(
                            f"صف {rno}: اسم الفرع «{vals.get('branch', '')}» يطابق أكثر من سجل أو اسم بديل؛ لا يمكن تحديد الفرع بالاسم وحده. أدخل كودًا صحيحًا أو عالج تعارض الأسماء أولًا."
                        )
                    elif branch_match_method == 'ambiguous_governorate':
                        errors.append(
                            f"صف {rno}: اسم المحافظة أو اسمها البديل يطابق أكثر من محافظة؛ صحح الاسم أو دليل المحافظات قبل الاستيراد."
                        )
                    elif branch_match_method == 'governorate_not_found':
                        errors.append(
                            f"صف {rno}: لم يُعثر على المحافظة أو كود فرع نشط مطابق: {vals.get('governorate', '')}."
                        )
                    else:
                        errors.append(
                            f"صف {rno}: لم يُعثر على فرع مطابق. راجع كود الفرع {vals.get('branch_code', '')} أو اسم الفرع {vals.get('branch', '')} والمحافظة {vals.get('governorate', '')}."
                        )
                    continue
                # Branch code determines the canonical branch and governorate.
                # Do not trust mismatched labels in Excel over a unique code.
                g = b.governorate
                required = ['name', 'job_title', 'job_code', 'hire_date']
                missing_values = [field for field in required if not vals.get(field)]
                if missing_values:
                    labels = {'name': 'اسم الموظف', 'job_title': 'الوظيفة', 'job_code': 'الكود الوظيفي', 'hire_date': 'تاريخ التعيين'}
                    errors.append(f"صف {rno}: البيانات الأساسية المطلوبة ناقصة: {'، '.join(labels[x] for x in missing_values)}.")
                    continue
                # Empty contact data is valid. Validate email only when supplied.
                vals['email'] = (vals.get('email') or '').strip()
                vals['company_phone'] = (vals.get('company_phone') or '').strip()
                vals['personal_phone'] = (vals.get('personal_phone') or '').strip()
                if vals['email'] and not valid_email(vals['email']):
                    errors.append(f'صف {rno}: البريد الإلكتروني غير صحيح.')
                    continue
                hd = _excel_date(_cell_by_index(row, mapping.get('hire_date')))
                if not hd:
                    errors.append(f'صف {rno}: تاريخ التعيين غير صحيح.')
                    continue
                normalized_code = _excel_employee_code_key(vals['job_code'])
                code_variants = _excel_employee_code_variants(normalized_code)
                ek = _excel_key(vals['email']) if vals['email'] else ''
                if code_variants & ambiguous_job_codes:
                    errors.append(f"صف {rno}: الكود الوظيفي يطابق أكثر من موظف بعد مراعاة بادئة 200 أو الأكواد البديلة؛ راجع الكود يدويًا.")
                    continue
                if code_variants & seen_job_codes:
                    errors.append(f"صف {rno}: الكود الوظيفي مكرر داخل ملف Excel: {vals.get('job_code', '')}.")
                    continue
                if ek and ek in seen_emails:
                    errors.append(f"صف {rno}: البريد الإلكتروني مكرر داخل ملف Excel: {vals.get('email', '')}.")
                    continue
                if ek and ek in ambiguous_emails:
                    errors.append(f"صف {rno}: البريد الإلكتروني مرتبط بأكثر من سجل في قاعدة البيانات؛ راجع البيانات يدويًا.")
                    continue
                seen_job_codes.update(code_variants)
                if ek:
                    seen_emails.add(ek)
                code_employee_ids = {
                    job_codes[variant] for variant in code_variants
                    if variant in job_codes and job_codes[variant] is not None
                }
                email_employee_id = emails.get(ek) if ek else None
                if len(code_employee_ids) > 1:
                    errors.append(f"صف {rno}: صيغ الكود الوظيفي تشير إلى أكثر من موظف؛ راجع الكود قبل الاستيراد.")
                    continue
                code_employee_id = next(iter(code_employee_ids), None)
                if code_employee_id is not None and email_employee_id is not None and code_employee_id != email_employee_id:
                    errors.append(f'صف {rno}: الكود الوظيفي والبريد الإلكتروني يعودان لموظفين مختلفين.')
                    continue
                matched_employee_id = code_employee_id if code_employee_id is not None else email_employee_id
                matched_identity = employee_identity_by_id.get(matched_employee_id) if matched_employee_id is not None else None
                if matched_identity and matched_identity.get('is_deleted'):
                    errors.append(
                        f"صف {rno}: تعارض يحتاج مراجعة: الكود الوظيفي أو البريد الإلكتروني يرتبط بموظف محذوف منطقيًا («{matched_identity.get('full_name') or 'موظف مسجل'}»). لم يُحدّث السجل ولم يُنشأ بديل تلقائيًا؛ راجع السجل قبل المتابعة."
                    )
                    continue
                # Email alone must not authorize changing an existing person's
                # personnel code. A code change is accepted as a match only
                # when the submitted code resolves to the same existing employee
                # through the canonical code or a confirmed alias/legacy variant.
                # Otherwise the row is a review conflict, not a routine update.
                if code_employee_id is None and email_employee_id is not None:
                    errors.append(
                        f"صف {rno}: تعارض يحتاج مراجعة: البريد الإلكتروني يطابق موظفًا قائمًا، لكن الكود الوظيفي في الملف لا يطابق كوده أو أي كود بديل مؤكد. لم يُربط الصف بالموظف ولم يُقترح تغيير الكود؛ صحح الكود أو راجع الهوية أولًا."
                    )
                    continue
                e = db.session.get(Employee, matched_employee_id) if matched_employee_id is not None else None
                import_name_key = (b.id, _excel_key(vals.get('name')))
                if not e and import_name_key in pending_employee_names:
                    errors.append(f"صف {rno}: الاسم نفسه تكرر لموظف جديد داخل الفرع في هذا الملف؛ راجع الكود والهوية قبل الإضافة.")
                    continue
                name_match_ids = employee_names_by_branch.get(import_name_key, set())
                if e and name_match_ids and e.id not in name_match_ids:
                    errors.append(f"صف {rno}: الكود والبريد يحددان سجلًا، لكن الاسم يطابق موظفًا آخر داخل الفرع؛ راجع الهوية قبل الاستيراد.")
                    continue
                if not e and name_match_ids:
                    if len(name_match_ids) == 1:
                        possible = employee_identity_by_id.get(next(iter(name_match_ids)))
                        errors.append(f"صف {rno}: الاسم يطابق موظفًا موجودًا في الفرع ({possible['full_name'] if possible else 'موظف مسجل'}) لكن الكود الوظيفي أو البريد لا يطابقه؛ صحح المعرف لتجنب إنشاء سجل مكرر.")
                    else:
                        errors.append(f"صف {rno}: الاسم يطابق أكثر من موظف داخل الفرع، والكود لا يحسم الهوية؛ راجع الكود والبريد يدويًا.")
                    continue
                if e:
                    previous_row = matched_employee_rows.get(e.id)
                    if previous_row is not None:
                        errors.append(
                            f"صف {rno}: تعارض في الهوية: هذا الموظف نفسه حُدد بالفعل في الصف {previous_row} من الملف؛ لا يمكن اعتماد صفين لتعديل السجل نفسه. راجع الكود الوظيفي والكود البديل."
                        )
                        continue
                    matched_employee_rows[e.id] = rno
                    changes = {}
                    candidate = {
                        'branch_id': b.id,
                        'branch': b.name,
                        'full_name': e.full_name if e.id in name_match_ids else vals['name'],
                        'email': vals['email'],
                        'job_title': vals['job_title'],
                        'job_code': vals['job_code'],
                        'hire_date': hd.isoformat(),
                        'company_phone': vals['company_phone'],
                        'personal_phone': vals['personal_phone'],
                    }
                    current = {
                        'branch_id': e.branch_id,
                        'branch': e.branch.name if e.branch else '—',
                        'full_name': e.full_name,
                        'email': e.email or '',
                        'job_title': e.job_title or '',
                        'job_code': e.job_code or '',
                        'hire_date': e.hire_date.isoformat() if e.hire_date else '',
                        'company_phone': e.company_phone or '',
                        'personal_phone': e.personal_phone or '',
                    }
                    for (f, label) in [
                        ('branch_id', 'الفرع'),
                        ('full_name', 'اسم الموظف'),
                        ('email', 'البريد الإلكتروني'),
                        ('job_title', 'الوظيفة'),
                        ('job_code', 'الكود الوظيفي'),
                        ('hire_date', 'تاريخ التعيين'),
                        ('company_phone', 'هاتف الشركة'),
                        ('personal_phone', 'الهاتف الشخصي'),
                    ]:
                        if str(current.get(f, '')) != str(candidate.get(f, '')):
                            changes[f] = {
                                'label': label,
                                'old': current.get(f, '—'),
                                'new': candidate.get(f, '—'),
                                # Keep IDs for validation/persistence, but show
                                # meaningful branch names during user approval.
                                'display_old': current.get('branch', '—') if f == 'branch_id' else current.get(f, '—'),
                                'display_new': candidate.get('branch', '—') if f == 'branch_id' else candidate.get(f, '—'),
                            }
                    if changes:
                        # Keep the identity evidence used by preview so approval
                        # can revalidate it. Selected fields alone are not enough:
                        # a code, confirmed alias, branch code, or corroborating
                        # email/name may have changed since the workbook was reviewed.
                        actions.append(
                            {
                                'type': 'employee',
                                'id': e.id,
                                'row': rno,
                                'name': e.full_name,
                                'changes': changes,
                                'branch_identity': {
                                    'branch_id': int(b.id),
                                    'governorate_id': int(b.governorate_id),
                                    'branch_code': vals.get('branch_code', ''),
                                    'branch_name': vals.get('branch', ''),
                                    'governorate': vals.get('governorate', ''),
                                    'match_method': branch_match_method,
                                },
                                'employee_identity': {
                                    'employee_id': int(e.id),
                                    # Source identity from Excel.
                                    'job_code': normalized_code,
                                    'email_key': ek,
                                    'email_matched': bool(ek and email_employee_id == e.id),
                                    'name_key': _excel_key(vals.get('name')),
                                    'name_matched': bool(name_match_ids == {e.id}),
                                    # Current identity values shown at preview time.
                                    # These are checked even when their fields are
                                    # not selected for mutation, so a changed target
                                    # cannot silently receive an unrelated update.
                                    'reviewed_branch_id': int(e.branch_id),
                                    'reviewed_full_name': e.full_name or '',
                                    'reviewed_email': e.email or '',
                                    'reviewed_job_code': e.job_code or '',
                                },
                            },
                        )
                    else:
                        actions.append({
                            'type': 'employee_match', 'id': e.id, 'row': rno,
                            'name': e.full_name, 'branch': b.name, 'governorate': g.name,
                        })
                else:
                    if any(variant in job_codes for variant in code_variants) or (ek and ek in emails):
                        errors.append(
                            f'صف {rno}: يوجد موظف مسجل مسبقًا بالكود الوظيفي أو البريد الإلكتروني.',
                        )
                        continue
                    actions.append(
                        {
                            'type': 'new_employee',
                            'row': rno,
                            'values': vals,
                            'hire_date': hd.isoformat(),
                            'branch_id': b.id,
                            'branch': b.name,
                            'governorate': g.name,
                            'branch_identity': {
                                'branch_id': int(b.id),
                                'governorate_id': int(b.governorate_id),
                                'branch_code': vals.get('branch_code', ''),
                                'branch_name': vals.get('branch', ''),
                                'governorate': vals.get('governorate', ''),
                                'match_method': branch_match_method,
                            },
                        },
                    )
                    for variant in code_variants:
                        job_codes[variant] = None
                    if ek:
                        emails[ek] = None
                    pending_employee_names.add(import_name_key)
        wb.close()
        review_path = _pending_excel_review_path(token)
        review_payload = json.dumps(
            {'kind': kind, 'mapping': mapping, 'actions': actions, 'errors': errors},
            ensure_ascii=False,
        ).encode('utf-8')
        write_private_bytes(review_path, review_payload)
        session['excel_review_token'] = token
        issue_labels = {
            'duplicate': 'مكرر / موجود مسبقًا',
            'conflict': 'تعارض يحتاج مراجعة',
            'unresolved': 'تعذر الربط الموثوق',
            'invalid': 'بيانات ناقصة أو غير صحيحة',
        }
        categorized_errors = [
            {'text': message, 'category': _excel_issue_category(message),
             'label': issue_labels[_excel_issue_category(message)]}
            for message in errors
        ]
        issue_counts = {category: sum(1 for item in categorized_errors if item['category'] == category)
                        for category in issue_labels}
        return render_template(
            'excel_import.html',
            kind=kind,
            review_step=True,
            actions=actions,
            errors=errors,
            categorized_errors=categorized_errors,
            issue_counts=issue_counts,
            fields=fields,
        )
    except Exception as exc:
        logger.error('Excel import confirmation failed (kind=%s, error=%s)', kind, type(exc).__name__)
        try:
            wb.close()
        except Exception:
            pass
        _remove_pending_excel(token)
        _clear_excel_import_session()
        flash('تعذر تحليل ملف Excel. لم يتم حفظ أي تغييرات؛ راجع الملف وحاول مرة أخرى.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))


@bp.post('/excel-import/<kind>/apply')
@req
def excel_import_apply(kind):
    _excel_import_permissions(kind)
    token = session.get('excel_review_token')
    path = _pending_excel_path(token) if token else None
    review_path = _pending_excel_review_path(token) if token else None
    if (
        not is_private_work_token(token)
        or session.get('excel_import_token') != token
        or session.get('excel_import_kind') != kind
        or not path
        or not os.path.isfile(path)
        or not review_path
        or not os.path.isfile(review_path)
        or pending_files_expired((path, review_path), _EXCEL_PENDING_TTL_SECONDS)
    ):
        _remove_pending_excel(token)
        _remove_pending_excel(session.get('excel_import_token'))
        _clear_excel_import_session()
        flash('انتهت جلسة مراجعة ملف Excel أو لم تعد متطابقة. ارفع الملف وراجعه مرة أخرى.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))
    if not _claim_pending_excel_apply(token):
        flash('عملية استيراد ملف Excel قيد التنفيذ أو سبق استهلاك رمزها؛ لم يتم تكرار التغييرات.')
        return redirect(url_for('excel_import.excel_import_page', kind=kind))
    try:
        touch_private_files((path, review_path))
        try:
            with open(review_path, 'r', encoding='utf-8') as fh:
                review = json.load(fh)
            if not isinstance(review, dict) or review.get('kind') != kind:
                _remove_pending_excel(token)
                _clear_excel_import_session()
                flash('نوع ملف المراجعة لا يطابق العملية المطلوبة. ارفع الملف وراجعه مرة أخرى.')
                return redirect(url_for('excel_import.excel_import_page', kind=kind))
            actions = review.get('actions', [])
            selection_plan = build_excel_import_selection_plan(
                kind, actions, request.form.getlist('change')
            )
            actions = list(selection_plan.actions)
            selected_adds = selection_plan.selected_adds
            selected_updates = selection_plan.selected_updates
            added = updated = 0
            skipped = selection_plan.skipped_adds
            errors = list(review.get('errors', []))

            # Revalidate the selected review manifest against current database state
            # before any writes. The service has no request/session-token ownership;
            # that remains in this blueprint.
            validate_selected_excel_import_changes(
                kind,
                selected_adds,
                selected_updates,
                selection_plan,
                session=db.session,
                Branch=Branch,
                Employee=Employee,
                Governorate=Governorate,
                TicketEmployeeCodeAlias=TicketEmployeeCodeAlias,
                EmployeeNameAlias=EmployeeNameAlias,
                excel_key=_excel_key,
                normalize_job_code=_excel_employee_code_key,
                job_code_lookup_variants=_excel_employee_code_variants,
                BranchNameAlias=BranchNameAlias,
                GovernorateNameAlias=GovernorateNameAlias,
            )
            counts = apply_selected_excel_import_changes(
                kind,
                actions,
                selection_plan,
                skipped,
                session=db.session,
                Governorate=Governorate,
                GovernorateNameAlias=GovernorateNameAlias,
                Branch=Branch,
                Employee=Employee,
                excel_key=_excel_key,
            )
            added, updated, skipped = counts.added, counts.updated, counts.skipped
            # Keep imported records and their audit entry in one transaction.
            log(
                'IMPORT',
                'Governorate' if kind == 'governorates' else ('Branch' if kind == 'branches' else 'Employee'),
                0,
                f'Excel: إضافة {added}، تحديث {updated}، تخطي دون تغييرات {skipped}',
            )
            db.session.commit()
            _remove_pending_excel(token)
            _clear_excel_import_session()
            return render_template(
                'excel_import.html',
                kind=kind,
                import_done=True,
                added=added,
                updated=updated,
                skipped=skipped,
                errors=errors,
            )
        except Exception as exc:
            db.session.rollback()
            logger.error('Excel import apply failed (kind=%s, error=%s)', kind, type(exc).__name__)
            _remove_pending_excel(token)
            _clear_excel_import_session()
            flash('تعذر اعتماد مراجعة Excel، أو تغيّرت البيانات منذ المعاينة. تم التراجع عن العملية ولم تُحفظ تغييرات؛ ارفع الملف وراجعه من جديد.')
            return redirect(url_for('excel_import.excel_import_page', kind=kind))
    finally:
        _release_pending_excel_apply(token)
