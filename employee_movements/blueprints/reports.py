"""Reports, exports and audit log."""

from flask import abort, Blueprint, render_template, request
from sqlalchemy import and_, case, func
from sqlalchemy.orm import joinedload
from datetime import date, timedelta

from ..access import can, req, roles
from ..constants import MOVEMENT_TYPES, STATUSES
from ..extensions import db
from ..models import Audit, Branch, Employee, Governorate, Movement, User, MissionEditRequest
from ..movement_state import assignment_closed_filter, assignment_open_filter
from ..validation import active_movement_types, parse_optional_iso_date
from ..services.export_safety import append_safe_xlsx_row, safe_csv_row

bp = Blueprint('reports', __name__)


def _apply_employee_type_report_filters(query, movement_type, branch_scope, args):
    """Apply identical, fail-closed filters to HTML, XLSX and CSV reports."""
    branch_scope = set(branch_scope or ())
    status = (args.get('status') or '').strip()
    employee_id = (args.get('employee_id') or '').strip()
    date_from = (args.get('date_from') or '').strip()
    date_to = (args.get('date_to') or '').strip()
    governorate_id = (args.get('governorate_id') or '').strip()
    branch_id = (args.get('branch_id') or '').strip()
    filter_error = False

    if status:
        if status in STATUSES:
            query = query.filter(Movement.status == status)
        else:
            query = query.filter(False)
            filter_error = True

    if employee_id:
        if employee_id.isdigit():
            # Base query already carries the authorized branch scope.
            query = query.filter(Movement.employee_id == int(employee_id))
        else:
            query = query.filter(False)
            filter_error = True

    # Report filters are global read filters; never derive their allowed
    # governorates from the selected operational work-scope. Mutation routes
    # perform their own selected-scope checks separately.
    scope_gov_ids = {
        governorate.id for governorate in Governorate.query.filter_by(is_active=True).all()
    }
    allowed_scope_branches = {
        branch.id: branch for branch in Branch.query.filter(
            Branch.is_active == True, Branch.id.in_(branch_scope or {-1})
        ).all()
        if branch.governorate_id in scope_gov_ids
    }
    selected_gov_id = None
    if governorate_id:
        if governorate_id.isdigit() and int(governorate_id) in scope_gov_ids:
            selected_gov_id = int(governorate_id)
            if not any(b.governorate_id == selected_gov_id for b in allowed_scope_branches.values()):
                query = query.filter(False)
                filter_error = True
            else:
                query = query.filter(Employee.branch.has(Branch.governorate_id == selected_gov_id))
        else:
            query = query.filter(False)
            filter_error = True

    if branch_id:
        if branch_id.isdigit() and int(branch_id) in allowed_scope_branches:
            selected_branch = allowed_scope_branches[int(branch_id)]
            if selected_gov_id is not None and selected_branch.governorate_id != selected_gov_id:
                query = query.filter(False)
                filter_error = True
            else:
                query = query.filter(Employee.branch_id == selected_branch.id)
        else:
            query = query.filter(False)
            filter_error = True

    date_from_obj = parse_optional_iso_date(date_from)
    date_to_obj = parse_optional_iso_date(date_to)
    if date_from and date_from_obj is None:
        query = query.filter(False)
        filter_error = True
    if date_to and date_to_obj is None:
        query = query.filter(False)
        filter_error = True
    date_range_error = bool(date_from_obj and date_to_obj and date_from_obj > date_to_obj)
    if date_range_error:
        # Invalid ranges fail closed. Never silently reverse the user's input;
        # the same helper is used by the HTML view and every spreadsheet export.
        query = query.filter(False)
        filter_error = True

    if date_from_obj:
        if movement_type == 'انتداب':
            # An open assignment has no end date, but it still overlaps the
            # requested period and must remain searchable.
            query = query.filter(db.or_(Movement.to_date.is_(None), Movement.to_date >= date_from_obj))
        elif movement_type == 'إجازة':
            query = query.filter(Movement.to_date >= date_from_obj)
        else:
            query = query.filter(Movement.permission_date >= date_from_obj)
    if date_to_obj:
        if movement_type in ('إجازة', 'انتداب'):
            query = query.filter(Movement.from_date <= date_to_obj)
        else:
            query = query.filter(Movement.permission_date <= date_to_obj)

    return query, {
        'status': status, 'employee_id': employee_id,
        'date_from': date_from, 'date_to': date_to,
        'date_from_obj': date_from_obj, 'date_to_obj': date_to_obj,
        'governorate_id': governorate_id, 'branch_id': branch_id,
        'filter_error': filter_error,
        'date_range_error': date_range_error,
    }


def _report_branch_scope():
    """Return branch scope for read-only reports, independent of work scope.

    All signed-in operational roles can read reports across active governorates.
    The work-governorate selector still limits any mutation action because the
    row-level authorization helpers (`branch_ok` / `can_manage_movement`) use
    the scoped endpoint. A report filter is a read filter, not an authorization
    grant and never changes the selected operational governorate.
    """
    raw_gid = (request.args.get('governorate_id') or '').strip()
    central = bool({'مسؤول التطبيق', 'Manager Application Support'} & roles())
    query = Branch.query.join(Governorate, Governorate.id == Branch.governorate_id).filter(
        Branch.is_active == True,
        Governorate.is_active == True,
    )
    if raw_gid:
        if not raw_gid.isdigit():
            abort(404)
        governorate = db.session.get(Governorate, int(raw_gid))
        if not governorate or not governorate.is_active:
            abort(404)
        query = query.filter(Branch.governorate_id == governorate.id)
    return {branch.id for branch in query.all()}, central


@bp.get('/manager')
@req
def manager_hub():
    """Shared cross-governorate reporting hub.

    All operational roles can read summaries. The post-close mission-edit request
    queue is deliberately exposed only to Manager Application Support, the sole
    reviewer/approver for those requests.
    """
    from ..access import roles
    rs = roles()
    if not can('view_reports'):
        abort(403)

    # Only branches under active governorates belong to the live organizational
    # scope. This also keeps summaries and pending-review queues consistent.
    governorates = Governorate.query.filter_by(is_active=True).order_by(Governorate.name.asc()).all()
    branches = (
        Branch.query.join(Governorate, Governorate.id == Branch.governorate_id)
        .filter(Branch.is_active == True, Governorate.is_active == True)
        .order_by(Branch.name.asc())
        .all()
    )
    allowed_branch_ids = {b.id for b in branches}
    is_manager = 'Manager Application Support' in rs
    pending_requests = []
    pending_count = 0
    if is_manager:
        pending_q = MissionEditRequest.query.filter_by(status='قيد المراجعة').join(Movement).join(
            Employee, Movement.employee_id == Employee.id
        )
        if allowed_branch_ids:
            pending_q = pending_q.filter(Employee.branch_id.in_(allowed_branch_ids))
        else:
            pending_q = pending_q.filter(False)
        pending_requests = pending_q.options(
            joinedload(MissionEditRequest.requester),
            joinedload(MissionEditRequest.movement).joinedload(Movement.employee),
        ).order_by(
            MissionEditRequest.created_at.asc(), MissionEditRequest.id.asc()
        ).limit(8).all()
        pending_count = pending_q.count()

    # Aggregate the whole manager summary in grouped SQL queries rather than
    # issuing five COUNT queries per governorate. The dashboard cost now grows
    # primarily with the number of returned governorates, not with query count.
    branch_counts = {governorate.id: 0 for governorate in governorates}
    for branch in branches:
        branch_counts[branch.governorate_id] = branch_counts.get(branch.governorate_id, 0) + 1

    employee_counts = {}
    movement_counts = {}
    if allowed_branch_ids:
        employee_counts = dict(
            db.session.query(Branch.governorate_id, func.count(Employee.id))
            .select_from(Employee)
            .join(Branch, Branch.id == Employee.branch_id)
            .filter(Employee.is_active == True, Branch.id.in_(allowed_branch_ids))
            .group_by(Branch.governorate_id)
            .all()
        )
        movement_rows = db.session.query(
            Branch.governorate_id.label('governorate_id'),
            func.sum(case((Movement.movement_type == 'إجازة', 1), else_=0)).label('leave_count'),
            func.sum(case((Movement.movement_type == 'انتداب', 1), else_=0)).label('assignment_count'),
            func.sum(case((Movement.movement_type == 'إذن', 1), else_=0)).label('permission_count'),
            func.sum(case((and_(Movement.movement_type == 'انتداب', assignment_closed_filter(Movement)), 1), else_=0)).label('closed_count'),
        ).select_from(Movement).join(
            Employee, Employee.id == Movement.employee_id
        ).join(
            Branch, Branch.id == Employee.branch_id
        ).filter(
            Movement.is_active == True,
            Branch.id.in_(allowed_branch_ids),
        ).group_by(Branch.governorate_id).all()
        movement_counts = {
            row.governorate_id: {
                'leave': int(row.leave_count or 0),
                'assignments': int(row.assignment_count or 0),
                'permissions': int(row.permission_count or 0),
                'closed_missions': int(row.closed_count or 0),
            }
            for row in movement_rows
        }

    governorate_reports = []
    for governorate in governorates:
        branch_count = branch_counts.get(governorate.id, 0)
        if not branch_count:
            continue
        aggregate = movement_counts.get(governorate.id, {})
        governorate_reports.append({
            'governorate': governorate,
            'branches': branch_count,
            'employees': int(employee_counts.get(governorate.id, 0)),
            'leave': aggregate.get('leave', 0),
            'assignments': aggregate.get('assignments', 0),
            'permissions': aggregate.get('permissions', 0),
            'closed_missions': aggregate.get('closed_missions', 0),
        })

    return render_template(
        'manager_hub.html',
        pending_requests=pending_requests,
        pending_count=pending_count,
        governorate_reports=governorate_reports,
        scope_has_data=bool(allowed_branch_ids),
        manager_hub_is_admin='مسؤول التطبيق' in rs,
        manager_hub_is_manager=is_manager,
    )


@bp.get('/audit')
@req
def audit():
    if not can('view_audit'):
        abort(403)
    rows = Audit.query.order_by(Audit.created_at.desc()).limit(1000).all()
    ids = {x.user_id for x in rows if x.user_id}
    audit_users = (
        {u.id: u.full_name for u in User.query.filter(User.id.in_(ids)).all()}
        if ids
        else {}
    )
    return render_template('audit.html', rows=rows, audit_users=audit_users)


@bp.get('/reports-missions')
@req
def reports_missions():
    if not (can('view_reports') or can('manage_movements')):
        abort(403)
    bs, is_central_manager = _report_branch_scope()
    requested_gid = (request.args.get('governorate_id') or '').strip()
    movement_q = (
        Movement.query.join(Employee).filter(Employee.branch_id.in_(bs), Movement.is_active == True)
        if bs else Movement.query.filter(False)
    )
    today = date.today()
    month_start = today.replace(day=1)
    next_month = (month_start.replace(day=28) + timedelta(days=4)).replace(day=1)
    week_end = today + timedelta(days=7)
    today_condition = (
        ((Movement.from_date <= today) & (Movement.to_date.is_(None) | (Movement.to_date >= today)))
        | (Movement.permission_date == today)
    )
    month_condition = (
        ((Movement.from_date >= month_start) & (Movement.from_date < next_month))
        | ((Movement.to_date >= month_start) & (Movement.to_date < next_month))
        | ((Movement.permission_date >= month_start) & (Movement.permission_date < next_month))
    )
    open_assignment_condition = and_(
        Movement.movement_type == 'انتداب',
        Movement.to_date.is_(None),
        assignment_open_filter(Movement),
    )
    follow_up_condition = and_(
        Movement.movement_type == 'انتداب',
        Movement.to_date.is_not(None),
        assignment_open_filter(Movement),
        Movement.to_date >= today,
        Movement.to_date <= week_end,
    )
    # One conditional aggregate replaces eight independent COUNT round-trips.
    aggregate_row = movement_q.order_by(None).with_entities(
        func.count(Movement.id).label('total'),
        func.coalesce(func.sum(case((Movement.movement_type == 'إجازة', 1), else_=0)), 0).label('leave_count'),
        func.coalesce(func.sum(case((Movement.movement_type == 'انتداب', 1), else_=0)), 0).label('assignment_count'),
        func.coalesce(func.sum(case((Movement.movement_type == 'إذن', 1), else_=0)), 0).label('permission_count'),
        func.coalesce(func.sum(case((open_assignment_condition, 1), else_=0)), 0).label('open_assignment_count'),
        func.coalesce(func.sum(case((today_condition, 1), else_=0)), 0).label('today_count'),
        func.coalesce(func.sum(case((month_condition, 1), else_=0)), 0).label('month_count'),
        func.coalesce(func.sum(case((follow_up_condition, 1), else_=0)), 0).label('follow_up_count'),
    ).one()
    stats = {
        'total': int(aggregate_row.total or 0),
        'leave': int(aggregate_row.leave_count or 0),
        'assignment': int(aggregate_row.assignment_count or 0),
        'permission': int(aggregate_row.permission_count or 0),
        'open_assignments': int(aggregate_row.open_assignment_count or 0),
        'today': int(aggregate_row.today_count or 0),
        'month': int(aggregate_row.month_count or 0),
        'follow_up_7d': int(aggregate_row.follow_up_count or 0),
    }
    selected_governorate = None
    raw_report_gid = (request.args.get('governorate_id') or '').strip()
    if raw_report_gid:
        if not raw_report_gid.isdigit():
            abort(404)
        selected_governorate = db.session.get(Governorate, int(raw_report_gid))
        if not selected_governorate or not selected_governorate.is_active:
            abort(404)
    return render_template(
        'reports_missions.html', report_stats=stats, report_today=today,
        report_week_end=week_end, selected_governorate=selected_governorate,
        manager_global_report=selected_governorate is None,
    )


@bp.get('/reports')
@req
def reports():
    if not can('view_reports'):
        abort(403)
    bs, _central_report = _report_branch_scope()
    q = (
        Movement.query.join(Employee).filter(Employee.branch_id.in_(bs), Movement.is_active == True)
        if bs
        else Movement.query.filter(False)
    )
    status = (request.args.get('status') or '').strip()
    mt = (request.args.get('movement_type') or '').strip()
    filter_error = False
    if status:
        if status in STATUSES:
            q = q.filter(Movement.status == status)
        else:
            q = q.filter(False)
            filter_error = True
    if mt:
        if mt in active_movement_types():
            q = q.filter(Movement.movement_type == mt)
        else:
            q = q.filter(False)
            filter_error = True
    summary = {status_value: 0 for status_value in STATUSES}
    for status_value, count in q.order_by(None).with_entities(
        Movement.status, func.count(Movement.id)
    ).group_by(Movement.status).all():
        if status_value in summary:
            summary[status_value] = int(count)

    by_type = {movement_type_value: 0 for movement_type_value in MOVEMENT_TYPES}
    for movement_type_value, count in q.order_by(None).with_entities(
        Movement.movement_type, func.count(Movement.id)
    ).group_by(Movement.movement_type).all():
        if movement_type_value in by_type:
            by_type[movement_type_value] = int(count)
    return render_template(
        'reports.html',
        summary=summary,
        by_type=by_type,
        status=status,
        movement_type=mt,
        statuses=STATUSES,
        movement_types=active_movement_types(),
        filter_error=filter_error,
    )


@bp.get('/reports/<report_type>')
@req
def employee_type_report(report_type):
    if not can('view_reports'):
        abort(403)
    mapping = {'leaves': 'إجازة', 'assignments': 'انتداب', 'permissions': 'إذن'}
    if report_type not in mapping:
        abort(404)
    mt = mapping[report_type]
    bs, _central_report = _report_branch_scope()
    q = (
        (
            Movement.query.join(Employee)
            .filter(
                Employee.branch_id.in_(bs),
                Movement.is_active == True,
                Movement.movement_type == mt,
            )
        )
        if bs
        else Movement.query.filter(False)
    )
    q, filters = _apply_employee_type_report_filters(q, mt, bs, request.args)
    status = filters['status']
    employee_id = filters['employee_id']
    date_from = filters['date_from']
    date_to = filters['date_to']
    filter_error = filters['filter_error']
    date_range_error = filters['date_range_error']
    rows = (
        q.order_by(
            Employee.full_name.asc(),
            Movement.from_date.desc(),
            Movement.permission_date.desc(),
            Movement.created_at.desc(),
        )
        .all()
    )
    grouped = []
    current = None
    for m in rows:
        if current is None or current['employee'].id != m.employee.id:
            current = {'employee': m.employee, 'rows': []}
            grouped.append(current)
        current['rows'].append(m)
    # Cascading report filters: governorate -> branch -> employee.
    selected_gov = request.args.get('governorate_id', '').strip()
    selected_branch = request.args.get('branch_id', '').strip()

    # Read filters can span all active governorates for every operational role.
    # Mutations triggered from these rows remain restricted by the selected work scope.
    allowed_gov_ids = {
        gov.id for gov in Governorate.query.filter_by(is_active=True).all()
    }
    allowed_govs = Governorate.query.filter(
        Governorate.is_active == True,
        Governorate.id.in_(allowed_gov_ids or {-1}),
    ).order_by(Governorate.name.asc()).all()
    allowed_branches = (
        Branch.query.filter(Branch.is_active == True, Branch.id.in_(bs or {-1}))
        .order_by(Branch.name.asc()).all()
    )
    if allowed_gov_ids:
        allowed_branches = [b for b in allowed_branches if b.governorate_id in allowed_gov_ids]
    else:
        allowed_branches = []

    selected_gov_obj = None
    if selected_gov.isdigit() and int(selected_gov) in allowed_gov_ids:
        selected_gov_obj = db.session.get(Governorate, int(selected_gov))
        allowed_branches = [b for b in allowed_branches if b.governorate_id == selected_gov_obj.id]
    else:
        selected_gov = ''

    selected_branch_obj = None
    if selected_branch.isdigit() and any((b.id == int(selected_branch) for b in allowed_branches)):
        selected_branch_obj = db.session.get(Branch, int(selected_branch))
    else:
        selected_branch = ''

    employees_q = (
        Employee.query.filter(Employee.is_active == True, Employee.branch_id.in_(bs))
        if bs
        else Employee.query.filter(False)
    )
    if selected_branch_obj:
        employees_q = employees_q.filter(Employee.branch_id == selected_branch_obj.id)
    elif selected_gov_obj:
        employees_q = employees_q.join(Branch).filter(Branch.governorate_id == selected_gov_obj.id)
    employees = employees_q.order_by(Employee.full_name.asc()).all()

    return render_template(
        'employee_type_report.html',
        report_type=report_type,
        title={'leaves': 'تقرير إجازات الموظفين', 'assignments': 'تقرير انتدابات الموظفين', 'permissions': 'تقرير أذونات الموظفين'}[report_type],
        rows=grouped,
        employees=employees,
        status=status,
        statuses=STATUSES,
        date_from=date_from,
        date_to=date_to,
        report_governorates=allowed_govs,
        report_branches=allowed_branches,
        selected_governorate=selected_gov,
        selected_branch=selected_branch,
        selected_employee=employee_id if employee_id.isdigit() else '',
        filter_error=filter_error,
        date_range_error=date_range_error,
    )


@bp.get('/reports/<report_type>.xlsx')
@req
def employee_type_report_xlsx(report_type):
    if not can('view_reports'):
        abort(403)
    mapping = {'leaves': 'إجازة', 'assignments': 'انتداب', 'permissions': 'إذن'}
    if report_type not in mapping:
        abort(404)
    mt = mapping[report_type]
    bs, _central_report = _report_branch_scope()
    q = (
        Movement.query.join(Employee).filter(
            Employee.branch_id.in_(bs), Movement.is_active == True, Movement.movement_type == mt
        ) if bs else Movement.query.filter(False)
    )
    q, _filters = _apply_employee_type_report_filters(q, mt, bs, request.args)
    rows = (
        q.order_by(
            Employee.full_name.asc(),
            Movement.from_date.desc(),
            Movement.permission_date.desc(),
        )
        .all()
    )
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from io import BytesIO
    wb = Workbook()
    ws = wb.active
    ws.title = 'التقرير'
    headers = [
        'الكود الوظيفي',
        'اسم الموظف',
        'المحافظة',
        'الفرع',
        'نوع الحركة',
        'نوع الإجازة',
        'جهة الانتداب',
        'من',
        'إلى',
        'تاريخ الإذن',
        'حالة الحركة',
    ]
    ws.append(headers)
    for m in rows:
        append_safe_xlsx_row(
            ws,
            [
                m.employee.job_code or '',
                m.employee.full_name,
                (
                    m.employee.branch.governorate.name
                    if m.employee.branch and m.employee.branch.governorate
                    else ''
                ),
                m.employee.branch.name if m.employee.branch else '',
                m.movement_type,
                m.leave_type or '',
                m.destination.name if m.destination else '',
                m.from_date or '',
                m.to_date or '',
                m.permission_date or '',
                m.status,
            ],
        )
    fill = PatternFill('solid', fgColor='1F3A5F')
    thin = Side(style='thin', color='E4E7EC')
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(horizontal='right', vertical='center')
            cell.border = Border(bottom=thin)
    widths = [16, 28, 18, 24, 16, 16, 24, 14, 14, 16, 14]
    for (i, w) in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    from flask import send_file
    return send_file(
        out,
        as_attachment=True,
        download_name=f'{report_type}_report.xlsx',
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )


@bp.get('/reports/<report_type>.csv')
@req
def employee_type_report_csv(report_type):
    if not can('view_reports'):
        abort(403)
    mapping = {'leaves': 'إجازة', 'assignments': 'انتداب', 'permissions': 'إذن'}
    if report_type not in mapping:
        abort(404)
    mt = mapping[report_type]
    import csv, io
    bs, _central_report = _report_branch_scope()
    q = (
        Movement.query.join(Employee).filter(
            Employee.branch_id.in_(bs), Movement.is_active == True, Movement.movement_type == mt
        ) if bs else Movement.query.filter(False)
    )
    q, _filters = _apply_employee_type_report_filters(q, mt, bs, request.args)
    rows = (
        q.order_by(
            Employee.full_name.asc(),
            Movement.from_date.desc(),
            Movement.permission_date.desc(),
        )
        .all()
    )
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(
        [
            'JobCode',
            'EmployeeName',
            'Governorate',
            'Branch',
            'MovementType',
            'LeaveType',
            'DestinationBranch',
            'From',
            'To',
            'PermissionDate',
            'Status',
            'RejectionReason',
        ],
    )
    for m in rows:
        w.writerow(
            safe_csv_row(
                [
                    m.employee.job_code or '',
                    m.employee.full_name,
                    (
                        m.employee.branch.governorate.name
                        if m.employee.branch and m.employee.branch.governorate
                        else ''
                    ),
                    m.employee.branch.name if m.employee.branch else '',
                    m.movement_type,
                    m.leave_type or '',
                    m.destination.name if m.destination else '',
                    m.from_date or '',
                    m.to_date or '',
                    m.permission_date or '',
                    m.status,
                    m.rejection_reason or '',
                ]
            ),
        )
    from flask import Response
    return Response(
        '\ufeff' + out.getvalue(),
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': f'attachment; filename={report_type}_employees_report.csv'},
    )


@bp.get('/reports.csv')
@req
def report_csv():
    if not can('view_reports'):
        abort(403)
    import csv, io
    bs, _central_report = _report_branch_scope()
    q = (
        Movement.query.join(Employee).filter(Employee.branch_id.in_(bs), Movement.is_active == True)
        if bs
        else Movement.query.filter(False)
    )
    status = (request.args.get('status') or '').strip()
    mt = (request.args.get('movement_type') or '').strip()
    if status:
        q = q.filter(Movement.status == status) if status in STATUSES else q.filter(False)
    if mt:
        q = q.filter(Movement.movement_type == mt) if mt in active_movement_types() else q.filter(False)
    rows = q.order_by(Movement.created_at.desc()).all()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(
        [
            'MovementID',
            'JobCode',
            'EmployeeName',
            'Branch',
            'MovementType',
            'LeaveType',
            'From',
            'To',
            'PermissionDate',
            'Status',
            'RejectionReason',
            'CreatedAt',
        ],
    )
    for m in rows:
        w.writerow(
            safe_csv_row(
                [
                    m.id,
                    m.employee.job_code or '',
                    m.employee.full_name,
                    m.employee.branch.name,
                    m.movement_type,
                    m.leave_type or '',
                    m.from_date or '',
                    m.to_date or '',
                    m.permission_date or '',
                    m.status,
                    m.rejection_reason or '',
                    m.created_at,
                ]
            ),
        )
    from flask import Response
    return Response(
        '\ufeff' + out.getvalue(),
        mimetype='text/csv; charset=utf-8',
        headers={'Content-Disposition': 'attachment; filename=movements_report.csv'},
    )
